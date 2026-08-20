import sqlite3
from datetime import date

from openpyxl import Workbook

from integracao_operacional_cobrancas import executar_integracao_planilha
from mapeamento_planilha_cobrancas import salvar_mapeamento
from migracao_modelo_dados import ensure_cobrancas_tables
from processador_agendado_cobrancas import ProcessadorAgendadoCobrancas


HEADERS = ['Órgão', 'Unidade', 'Fornecedor', 'Serviço', 'Competência', 'Vencimento', 'Valor original', 'Valor pago', 'Valor pendente', 'Responsável']
MAPEAMENTO = {
    'orgao': 'Órgão', 'unidade': 'Unidade', 'fornecedor': 'Fornecedor', 'servico': 'Serviço',
    'competencia': 'Competência', 'data_vencimento': 'Vencimento', 'valor_original': 'Valor original',
    'valor_pago': 'Valor pago', 'valor_pendente': 'Valor pendente', 'responsavel': 'Responsável',
}


def _ambiente(tmp_path, rows=None, headers=HEADERS, aba='Oficial', mapping_headers=None):
    tmp_path.mkdir(parents=True, exist_ok=True)
    planilha = tmp_path / 'cobrancas.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = aba
    sheet.append(headers)
    sheet.append(rows or ['SRA', 'Sede', 'Fornecedor', 'Servico', '08/2026', '24/08/2026', 100, 0, 100, 'Ana'])
    workbook.save(planilha)
    workbook.close()
    mapeamento = tmp_path / 'mapeamento.json'
    salvar_mapeamento(mapeamento, MAPEAMENTO, mapping_headers or headers, 'teste', aba)
    config = {
        'cobrancas_planilha_oficial_path': str(planilha), 'cobrancas_planilha_aba': aba,
        'cobrancas_mapeamento_path': str(mapeamento), 'cobrancas_dias_antecedencia': '5',
        'cobrancas_dias_vencida': '0',
    }
    conn = sqlite3.connect(':memory:')
    conn.row_factory = sqlite3.Row
    conn.execute('CREATE TABLE faturas (id INTEGER PRIMARY KEY)')
    ensure_cobrancas_tables(conn)
    return conn, config, mapeamento


def test_integracao_aplica_mapeamento_e_atualiza_sem_duplicidade(tmp_path):
    conn, config, _ = _ambiente(tmp_path)
    primeira = executar_integracao_planilha(conn, config, data_atual=date(2026, 8, 20))
    segunda = executar_integracao_planilha(conn, config, data_atual=date(2026, 8, 20))
    assert primeira['status'] == 'sucesso'
    assert primeira['mapeamento_versao'] == 1
    assert primeira['criadas'] == 1
    assert segunda['atualizadas'] == 1
    assert conn.execute('SELECT COUNT(*) FROM contas').fetchone()[0] == 1
    assert conn.execute('SELECT status FROM contas').fetchone()[0] == 'PROXIMA_DO_VENCIMENTO'


def test_integracao_registra_hash_e_mudancas_de_valor_e_vencimento(tmp_path):
    conn, config, _ = _ambiente(tmp_path)
    executar_integracao_planilha(conn, config, data_atual=date(2026, 8, 20))
    planilha = config['cobrancas_planilha_oficial_path']
    workbook = __import__('openpyxl').load_workbook(planilha)
    sheet = workbook['Oficial']
    sheet['G2'] = 125
    sheet['I2'] = 125
    sheet['F2'] = '15/09/2026'
    workbook.save(planilha)
    workbook.close()
    resultado = executar_integracao_planilha(conn, config, data_atual=date(2026, 8, 20))
    conta = conn.execute('SELECT valor_pendente, data_vencimento FROM contas').fetchone()
    assert resultado['atualizadas'] == 1
    assert conta['valor_pendente'] == 125
    assert conta['data_vencimento'] == '2026-09-15'
    assert resultado['divergencias']
    assert conn.execute('SELECT COUNT(*) FROM leituras_planilha').fetchone()[0] == 2


def test_alteracao_do_mapeamento_ativo_e_aplicada(tmp_path):
    conn, config, caminho_mapeamento = _ambiente(tmp_path)
    workbook = __import__('openpyxl').load_workbook(config['cobrancas_planilha_oficial_path'])
    workbook['Oficial']['C1'] = 'Prestador'
    workbook.save(config['cobrancas_planilha_oficial_path'])
    workbook.close()
    novo = dict(MAPEAMENTO)
    novo['fornecedor'] = 'Prestador'
    headers_alterados = ['Prestador' if header == 'Fornecedor' else header for header in HEADERS]
    salvar_mapeamento(caminho_mapeamento, novo, headers_alterados, 'teste-2', 'Oficial')
    resultado = executar_integracao_planilha(conn, config, data_atual=date(2026, 8, 20))
    assert resultado['status'] == 'sucesso'
    assert resultado['mapeamento_versao'] == 2


def test_mapeamento_ausente_ou_aba_inexistente_interrompe_em_revisao(tmp_path):
    conn, config, mapeamento = _ambiente(tmp_path, aba='Oficial')
    mapeamento.unlink()
    resultado = executar_integracao_planilha(conn, config)
    assert resultado['status'] == 'AGUARDANDO_REVISÃO'
    conn.close()

    conn, config, _ = _ambiente(tmp_path / 'outro', aba='Oficial')
    config['cobrancas_planilha_aba'] = 'Ausente'
    resultado = executar_integracao_planilha(conn, config)
    assert resultado['status'] == 'AGUARDANDO_REVISÃO'


def test_coluna_obrigatoria_ausente_interrompe_em_revisao(tmp_path):
    headers = [item for item in HEADERS if item != 'Valor pendente']
    conn, config, _ = _ambiente(tmp_path, headers=headers, mapping_headers=HEADERS, rows=['SRA', 'Sede', 'Fornecedor', 'Servico', '08/2026', '31/08/2026', 100, 0, 'Ana'])
    resultado = executar_integracao_planilha(conn, config)
    assert resultado['status'] == 'AGUARDANDO_REVISÃO'


def test_api_manual_usa_o_leitor_e_o_mapeamento_ativos(monkeypatch, tmp_path):
    import sys
    sys.path.insert(0, 'Interface_Web/backend')
    import app as app_module

    conn, config, _ = _ambiente(tmp_path)
    monkeypatch.setattr(app_module, 'get_db', lambda: conn)
    monkeypatch.setattr(app_module, 'load_config', lambda: config)
    monkeypatch.setattr(app_module, 'usuario_da_requisicao', lambda _conn: {'usuario': 'teste', 'perfil': 'administrador'})
    cliente = app_module.app.test_client()
    resposta = cliente.post('/api/cobrancas/planilha/ler')
    assert resposta.status_code == 200
    assert resposta.get_json()['mapeamento_versao'] == 1


def test_agendador_usa_a_mesma_integracao_real(tmp_path):
    conn, config, _ = _ambiente(tmp_path)
    processador = ProcessadorAgendadoCobrancas(
        config, conn=conn, raiz=str(tmp_path),
        agora=lambda: __import__('datetime').datetime(2026, 8, 20, 12, 0, tzinfo=__import__('datetime').timezone.utc),
        registrar=lambda resultado: None,
        ler_planilha=lambda: (_ for _ in ()).throw(AssertionError('leitor injetado não deveria ser usado')),
        atualizar_contas=lambda *_: {}, recalcular_status=lambda *_: [], aplicar_regras=lambda *_: [],
        gerar_previa=lambda *_: None, enviar=lambda *_: None, verificar_duplicidade=lambda *_: False,
    )
    resultado = processador.executar()
    assert resultado['status'] == 'concluido'
    assert conn.execute('SELECT COUNT(*) FROM contas').fetchone()[0] == 1