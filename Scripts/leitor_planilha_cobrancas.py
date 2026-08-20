"""Leitura somente de consulta da planilha oficial de cobrancas."""

from __future__ import annotations

import hashlib
import re
import unicodedata
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Optional

from openpyxl import load_workbook


CAMPOS = (
    'orgao',
    'unidade',
    'fornecedor',
    'servico',
    'numero_fatura',
    'competencia',
    'data_emissao',
    'data_vencimento',
    'valor_original',
    'valor_rateado',
    'valor_pago',
    'valor_pendente',
    'situacao',
    'observacao',
    'telefone_destinatario',
    'responsavel',
)

CAMPOS_OBRIGATORIOS = (
    'orgao',
    'fornecedor',
    'servico',
    'competencia',
    'data_vencimento',
    'valor_original',
    'valor_pago',
    'valor_pendente',
)

ALIASES = {
    'orgao': ('orgao', 'órgão', 'orgao devedor'),
    'unidade': ('unidade', 'unidade beneficiaria', 'unidade beneficiária'),
    'fornecedor': ('fornecedor', 'prestador'),
    'servico': ('servico', 'serviço', 'descricao do servico', 'descrição do serviço'),
    'numero_fatura': ('numero da fatura', 'número da fatura', 'fatura', 'identificador'),
    'competencia': ('competencia', 'competência', 'mes', 'mês'),
    'data_emissao': ('data de emissao', 'data de emissão', 'emissao', 'emissão'),
    'data_vencimento': ('data de vencimento', 'vencimento', 'data vencimento'),
    'valor_original': ('valor original', 'valor bruto', 'valor total', 'total'),
    'valor_rateado': ('valor rateado', 'rateio', 'valor do rateio'),
    'valor_pago': ('valor pago', 'pago', 'pagamento'),
    'valor_pendente': ('valor pendente', 'pendente', 'saldo pendente', 'saldo'),
    'situacao': ('situacao', 'situação', 'status'),
    'observacao': ('observacao', 'observação', 'observacoes', 'observações'),
    'telefone_destinatario': ('telefone', 'telefone destinatario', 'telefone destinatário'),
    'responsavel': ('responsavel', 'responsável', 'responsavel pelo pagamento', 'responsável pelo pagamento'),
}


class ErroLeituraPlanilha(ValueError):
    """Erro detalhado e seguro da leitura da planilha."""

    def __init__(self, codigo: str, mensagem: str, **detalhes: Any) -> None:
        super().__init__(mensagem)
        self.codigo = codigo
        self.mensagem = mensagem
        self.detalhes = detalhes

    def as_dict(self) -> Dict[str, Any]:
        return {'codigo': self.codigo, 'mensagem': self.mensagem, **self.detalhes}


class LeitorPlanilhaCobrancas:
    def __init__(
        self,
        config: Dict[str, Any],
        workbook_loader: Callable[..., Any] = load_workbook,
        clock: Callable[[], datetime] | None = None,
    ) -> None:
        self.config = config
        self.workbook_loader = workbook_loader
        self.clock = clock or (lambda: datetime.now(timezone.utc))

    def ler(
        self,
        mapeamento: Optional[Dict[str, str]] = None,
        limite_previa: int = 10,
    ) -> Dict[str, Any]:
        caminho = self._caminho_configurado()
        self._validar_arquivo(caminho)
        try:
            hash_origem = self._hash(caminho)
        except (PermissionError, OSError) as exc:
            raise ErroLeituraPlanilha('arquivo_bloqueado', f'Planilha inacessível ou bloqueada durante a leitura: {exc}') from exc
        lida_em = self.clock().astimezone(timezone.utc).isoformat()
        try:
            workbook = self.workbook_loader(caminho, read_only=True, data_only=False)
        except (PermissionError, OSError) as exc:
            raise ErroLeituraPlanilha('arquivo_bloqueado', f'Planilha inacessível ou bloqueada: {exc}') from exc
        except Exception as exc:
            raise ErroLeituraPlanilha('arquivo_invalido', f'Não foi possível abrir a planilha: {exc}') from exc

        try:
            abas = list(workbook.sheetnames)
            aba_nome = str(self.config.get('cobrancas_planilha_aba', '')).strip()
            if not aba_nome:
                raise ErroLeituraPlanilha('aba_nao_configurada', 'Aba da planilha de cobranças não configurada.', abas=abas)
            if aba_nome not in abas:
                raise ErroLeituraPlanilha('aba_inexistente', f'Aba configurada não encontrada: {aba_nome}.', abas=abas)
            planilha = workbook[aba_nome]
            cabecalhos, linha_cabecalho = self._identificar_cabecalhos(planilha, mapeamento)
            if not cabecalhos:
                raise ErroLeituraPlanilha('cabecalhos_nao_reconhecidos', 'Nenhum cabeçalho reconhecível foi encontrado.', aba=aba_nome)
            ausentes = [campo for campo in CAMPOS_OBRIGATORIOS if campo not in cabecalhos]
            if ausentes:
                raise ErroLeituraPlanilha(
                    'colunas_ausentes',
                    'Colunas obrigatórias ausentes: ' + ', '.join(ausentes) + '.',
                    aba=aba_nome,
                    cabecalhos_encontrados=list(cabecalhos.values()),
                    campos_ausentes=ausentes,
                )
            linhas, invalidas, vazias, previa = self._ler_linhas(planilha, cabecalhos, linha_cabecalho, limite_previa)
            return {
                'status': 'sucesso',
                'arquivo': str(caminho),
                'aba': aba_nome,
                'abas': abas,
                'hash_origem': hash_origem,
                'versao_origem': hash_origem,
                'lida_em': lida_em,
                'linha_cabecalho': linha_cabecalho,
                'cabecalhos': cabecalhos,
                'linhas': linhas,
                'linhas_invalidas': invalidas,
                'celulas_vazias': vazias,
                'previa': previa,
            }
        finally:
            workbook.close()

    def inspecionar(self, limite_previa: int = 10) -> Dict[str, Any]:
        """Lista abas, cabecalhos e linhas brutas para montar um mapeamento."""
        caminho = self._caminho_configurado()
        self._validar_arquivo(caminho)
        try:
            hash_origem = self._hash(caminho)
            workbook = self.workbook_loader(caminho, read_only=True, data_only=False)
        except (PermissionError, OSError) as exc:
            raise ErroLeituraPlanilha('arquivo_bloqueado', f'Planilha inacessível ou bloqueada: {exc}') from exc
        except Exception as exc:
            raise ErroLeituraPlanilha('arquivo_invalido', f'Não foi possível abrir a planilha: {exc}') from exc
        try:
            abas = list(workbook.sheetnames)
            aba_nome = str(self.config.get('cobrancas_planilha_aba', '')).strip()
            if not aba_nome:
                raise ErroLeituraPlanilha('aba_nao_configurada', 'Aba da planilha de cobranças não configurada.', abas=abas)
            if aba_nome not in abas:
                raise ErroLeituraPlanilha('aba_inexistente', f'Aba configurada não encontrada: {aba_nome}.', abas=abas)
            planilha = workbook[aba_nome]
            cabecalhos = self._cabecalhos_brutos(planilha)
            previa = []
            for numero_linha, valores in enumerate(planilha.iter_rows(min_row=cabecalhos['linha'] + 1, max_row=cabecalhos['linha'] + limite_previa, values_only=True), cabecalhos['linha'] + 1):
                if any(valor is not None and str(valor).strip() for valor in valores):
                    previa.append({'linha': numero_linha, 'valores': list(valores)})
            return {
                'status': 'sucesso',
                'arquivo': str(caminho),
                'aba': aba_nome,
                'abas': abas,
                'hash_origem': hash_origem,
                'lida_em': self.clock().astimezone(timezone.utc).isoformat(),
                'cabecalhos': cabecalhos['nomes'],
                'linha_cabecalho': cabecalhos['linha'],
                'previa': previa,
            }
        finally:
            workbook.close()

    def _caminho_configurado(self) -> Path:
        valor = str(self.config.get('cobrancas_planilha_oficial_path', '')).strip()
        if not valor:
            raise ErroLeituraPlanilha('arquivo_nao_configurado', 'Caminho da planilha oficial de cobranças não configurado.')
        return Path(valor)

    @staticmethod
    def _validar_arquivo(caminho: Path) -> None:
        if not caminho.exists():
            raise ErroLeituraPlanilha('arquivo_inexistente', f'Planilha oficial não encontrada: {caminho}')
        if not caminho.is_file():
            raise ErroLeituraPlanilha('arquivo_invalido', f'O caminho configurado não é um arquivo: {caminho}')
        try:
            with caminho.open('rb'):
                pass
        except (PermissionError, OSError) as exc:
            raise ErroLeituraPlanilha('arquivo_bloqueado', f'Planilha inacessível ou bloqueada: {exc}') from exc

    @staticmethod
    def _hash(caminho: Path) -> str:
        digest = hashlib.sha256()
        with caminho.open('rb') as arquivo:
            for bloco in iter(lambda: arquivo.read(1024 * 1024), b''):
                digest.update(bloco)
        return digest.hexdigest()

    def _identificar_cabecalhos(self, planilha: Any, mapeamento: Optional[Dict[str, str]]) -> tuple[Dict[str, str], int]:
        for numero_linha, linha in enumerate(planilha.iter_rows(min_row=1, max_row=min(planilha.max_row, 20), values_only=True), 1):
            encontrados: Dict[str, str] = {}
            valores = {_normalizar(valor): str(valor).strip() for valor in linha if valor is not None and str(valor).strip()}
            if mapeamento:
                for campo, cabecalho in mapeamento.items():
                    if _normalizar(cabecalho) in valores:
                        encontrados[campo] = valores[_normalizar(cabecalho)]
            else:
                for campo, aliases in ALIASES.items():
                    for alias in aliases:
                        if _normalizar(alias) in valores:
                            encontrados[campo] = valores[_normalizar(alias)]
                            break
            if encontrados:
                return encontrados, numero_linha
        return {}, 0

    @staticmethod
    def _cabecalhos_brutos(planilha: Any) -> Dict[str, Any]:
        for numero_linha, linha in enumerate(planilha.iter_rows(min_row=1, max_row=min(planilha.max_row, 20), values_only=True), 1):
            nomes = [str(valor).strip() for valor in linha if valor is not None and str(valor).strip()]
            if nomes:
                return {'linha': numero_linha, 'nomes': nomes}
        raise ErroLeituraPlanilha('cabecalhos_nao_reconhecidos', 'Nenhum cabeçalho reconhecível foi encontrado.')

    def _ler_linhas(self, planilha: Any, cabecalhos: Dict[str, str], linha_cabecalho: int, limite_previa: int) -> tuple[list[dict], list[dict], list[dict], list[dict]]:
        indices = {str(valor).strip(): indice for indice, valor in enumerate(next(planilha.iter_rows(min_row=linha_cabecalho, max_row=linha_cabecalho, values_only=True)), 1) if valor is not None}
        linhas: list[dict] = []
        invalidas: list[dict] = []
        vazias: list[dict] = []
        previa: list[dict] = []
        for numero_linha, valores in enumerate(planilha.iter_rows(min_row=linha_cabecalho + 1, values_only=True), linha_cabecalho + 1):
            if not any(valor is not None and str(valor).strip() for valor in valores):
                continue
            registro: Dict[str, Any] = {'linha': numero_linha}
            erros: list[str] = []
            vazios_linha: list[str] = []
            for campo, cabecalho in cabecalhos.items():
                valor_bruto = valores[indices[cabecalho] - 1] if indices.get(cabecalho, 0) <= len(valores) else None
                if valor_bruto is None or str(valor_bruto).strip() == '':
                    registro[campo] = None
                    vazios_linha.append(campo)
                    continue
                if campo in {'valor_original', 'valor_rateado', 'valor_pago', 'valor_pendente'}:
                    valor, erro = _moeda(valor_bruto)
                elif campo in {'data_emissao', 'data_vencimento'}:
                    valor, erro = _data_brasileira(valor_bruto)
                else:
                    valor, erro = _texto(valor_bruto), None
                registro[campo] = valor
                if erro:
                    erros.append(f'{campo}: {erro}')
            if vazios_linha:
                vazias.append({'linha': numero_linha, 'campos': vazios_linha})
            if erros:
                invalidas.append({'linha': numero_linha, 'erros': erros, 'valores': registro})
            else:
                linhas.append(registro)
            if len(previa) < max(0, limite_previa):
                previa.append(registro.copy())
        return linhas, invalidas, vazias, previa


def _normalizar(valor: Any) -> str:
    texto = unicodedata.normalize('NFKD', str(valor or '')).encode('ascii', 'ignore').decode('ascii')
    texto = re.sub(r'[^A-Z0-9]+', ' ', texto.upper())
    return ' '.join(texto.split())


def _texto(valor: Any) -> str:
    return str(valor).strip()


def _moeda(valor: Any) -> tuple[Optional[float], Optional[str]]:
    if isinstance(valor, bool):
        return None, 'valor booleano inválido'
    if isinstance(valor, (int, float)):
        return float(valor), None
    texto = str(valor).strip().upper()
    if not texto or texto in {'N/A', 'NA', '-', '--'}:
        return None, 'valor não informado'
    negativo = texto.startswith('(') and texto.endswith(')')
    texto = texto.replace('R$', '').replace(' ', '').replace('.', '').replace(',', '.')
    if negativo:
        texto = '-' + texto[1:-1]
    try:
        return float(texto), None
    except ValueError:
        return None, f'valor monetário inválido: {valor}'


def _data_brasileira(valor: Any) -> tuple[Optional[str], Optional[str]]:
    if isinstance(valor, datetime):
        return valor.date().isoformat(), None
    if isinstance(valor, date):
        return valor.isoformat(), None
    texto = str(valor).strip()
    for formato in ('%d/%m/%Y', '%d/%m/%y'):
        try:
            return datetime.strptime(texto, formato).date().isoformat(), None
        except ValueError:
            continue
    return None, f'data brasileira inválida: {valor}'
