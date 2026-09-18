"""Integração segura de resultados calculados com a planilha oficial."""

from __future__ import annotations

import json
import gc
import os
import shutil
import tempfile
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Dict, Iterable, Optional

from openpyxl import load_workbook
from backup_recuperacao import criar_backup


CAMPOS_CHAVE = ('servico', 'fornecedor', 'complemento', 'secao', 'competencia')
ALIASES = {
    'servico': ('servico', 'serviço'),
    'fornecedor': ('fornecedor',),
    'complemento': ('complemento',),
    'secao': ('secao', 'seção'),
    'competencia': ('competencia', 'competência'),
}


class IntegradorPlanilha:
    def __init__(self, config: Dict[str, Any], usuario: str = 'sistema', audit_sink: Optional[Callable[[Dict[str, Any]], None]] = None):
        self.config = config
        self.usuario = usuario
        self.audit_sink = audit_sink

    def gravar(self, fatura: Dict[str, Any], calculo: Dict[str, Any]) -> Dict[str, Any]:
        validacao = self._validar_precondicoes()
        if validacao:
            return self._revisao(validacao)
        if calculo.get('status') != 'CALCULADO_VALIDADO':
            return self._revisao('Resultado de cálculo não está validado.')

        caminho = Path(self.config['planilha_oficial_path'])
        backup = self._criar_backup(caminho)
        if isinstance(backup, str):
            return self._revisao(backup)

        temp_path = None
        try:
            with self._arquivo_desbloqueado(caminho):
                pass
            workbook = load_workbook(caminho, keep_vba=caminho.suffix.lower() == '.xlsm')
            planilha = workbook[self.config['planilha_aba']] if self.config.get('planilha_aba') else workbook.active
            cabecalhos, linha_cabecalho = self._localizar_cabecalhos(planilha)
            if not cabecalhos:
                return self._revisao('Cabeçalhos obrigatórios não encontrados na planilha.')
            linha = self._localizar_linha(planilha, linha_cabecalho, cabecalhos, fatura)
            if linha is None:
                return self._revisao('Correspondência da fatura não encontrada; nenhuma linha foi criada.')

            parcelas = calculo.get('parcelas') or []
            colunas_orgao = self._localizar_colunas_orgaos(planilha, linha_cabecalho, parcelas)
            if isinstance(colunas_orgao, str):
                return self._revisao(colunas_orgao)

            auditoria = []
            for parcela in parcelas:
                coluna = colunas_orgao[parcela['orgao']]
                celula = planilha.cell(row=linha, column=coluna)
                anterior = celula.value
                novo = float(parcela['valor'])
                celula.value = novo
                auditoria.append(self._auditoria(celula.coordinate, anterior, novo, caminho, calculo))

            temp_handle = tempfile.NamedTemporaryFile(delete=False, suffix=caminho.suffix, dir=caminho.parent)
            temp_path = Path(temp_handle.name)
            temp_handle.close()
            workbook.save(temp_path)
            workbook.close()
            del workbook
            gc.collect()
            os.replace(temp_path, caminho)
            for registro in auditoria:
                self._registrar_auditoria(registro)
            self._registrar_auditoria({
                'tipo_evento': 'backup', 'arquivo': str(caminho),
                'pacote': str(backup), 'usuario': self.usuario,
                'data': datetime.now(timezone.utc).isoformat(), 'status': 'sucesso',
            })
            return {
                'status': 'GRAVADO_VALIDADO',
                'planilha': str(caminho),
                'backup': str(backup),
                'linha': linha,
                'celulas': auditoria,
                'regra': calculo.get('regra'),
            }
        except (OSError, PermissionError, ValueError, KeyError) as exc:
            if temp_path and temp_path.exists():
                temp_path.unlink(missing_ok=True)
            return self._revisao(f'Falha segura na gravação: {exc}')

    def _validar_precondicoes(self) -> Optional[str]:
        caminho = str(self.config.get('planilha_oficial_path', '')).strip()
        if not caminho:
            return 'Caminho da planilha oficial não configurado.'
        arquivo = Path(caminho)
        if not arquivo.is_file():
            return 'Planilha oficial não existe.'
        backup_path = str(self.config.get('backup_path', '')).strip()
        if not backup_path:
            return 'Caminho de backup não configurado.'
        if str(self.config.get('ambiente', '')).strip().lower() == 'producao':
            return 'Gravação bloqueada em ambiente de produção; use uma cópia de teste.'
        try:
            with arquivo.open('rb'):
                pass
        except OSError as exc:
            return f'Planilha bloqueada ou inacessível: {exc}'
        return None

    def _criar_backup(self, caminho: Path) -> Path | str:
        pasta = Path(self.config['backup_path'])
        try:
            pasta.mkdir(parents=True, exist_ok=True)
            pacote = criar_backup(self.config, motivo='alteracao_planilha', versao=str(self.config.get('versao_sistema', 'v1.0.0')), destino=pasta)
            destino = pasta / f'{caminho.stem}_{datetime.now().strftime("%Y%m%d_%H%M%S_%f")}{caminho.suffix}'
            shutil.copy2(caminho, destino)
            return pacote
        except OSError as exc:
            return f'Não foi possível criar backup: {exc}'

    def _localizar_cabecalhos(self, planilha):
        for numero_linha in range(1, min(planilha.max_row, 20) + 1):
            encontrados = {}
            for celula in planilha[numero_linha]:
                valor = _normalizar(celula.value)
                for campo, aliases in ALIASES.items():
                    if valor in {_normalizar(alias) for alias in aliases}:
                        encontrados[campo] = celula.column
            if all(campo in encontrados for campo in CAMPOS_CHAVE):
                return encontrados, numero_linha
        return None, None

    def _localizar_linha(self, planilha, linha_cabecalho, cabecalhos, fatura):
        esperados = {campo: _normalizar(fatura.get(campo)) for campo in CAMPOS_CHAVE}
        if any(not valor for valor in esperados.values()):
            return None
        for numero_linha in range(linha_cabecalho + 1, planilha.max_row + 1):
            if all(_normalizar(planilha.cell(numero_linha, cabecalhos[campo]).value) == esperado for campo, esperado in esperados.items()):
                return numero_linha
        return None

    def _localizar_colunas_orgaos(self, planilha, linha_cabecalho, parcelas):
        resultado = {}
        for celula in planilha[linha_cabecalho]:
            cabecalho = _normalizar(celula.value)
            for parcela in parcelas:
                orgao = _normalizar(parcela.get('orgao'))
                if cabecalho == orgao or cabecalho == f'{orgao} R$':
                    resultado[parcela['orgao']] = celula.column
        faltantes = [parcela['orgao'] for parcela in parcelas if parcela['orgao'] not in resultado]
        return f'Coluna de órgão ausente: {faltantes[0]}.' if faltantes else resultado

    def _auditoria(self, celula, anterior, novo, caminho, calculo):
        return {
            'celula': celula,
            'valor_anterior': anterior,
            'valor_novo': novo,
            'planilha': str(caminho),
            'usuario': self.usuario,
            'data': datetime.now(timezone.utc).isoformat(),
            'regra': calculo.get('regra'),
        }

    def _registrar_auditoria(self, registro):
        if self.audit_sink:
            self.audit_sink(registro)

    @staticmethod
    def _revisao(motivo):
        return {'status': 'AGUARDANDO_REVISÃO', 'motivos_revisao': [motivo], 'celulas': []}

    @staticmethod
    def _arquivo_desbloqueado(caminho):
        return _ArquivoDesbloqueado(caminho)


class _ArquivoDesbloqueado:
    def __init__(self, caminho):
        self.caminho = caminho
        self.handle = None

    def __enter__(self):
        try:
            self.handle = self.caminho.open('r+b')
            return self
        except OSError as exc:
            self.__exit__(None, None, None)
            raise PermissionError(f'Planilha bloqueada: {exc}') from exc

    def __exit__(self, exc_type, exc_value, traceback):
        if self.handle:
            self.handle.close()


def _normalizar(valor: Any) -> str:
    if valor is None:
        return ''
    texto = unicodedata.normalize('NFKD', str(valor)).encode('ascii', 'ignore').decode('ascii')
    return ' '.join(texto.strip().upper().split())