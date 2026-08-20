from pathlib import Path
import sqlite3

from openpyxl import Workbook, load_workbook

from calculador_rateio import CalculadorRateioCompartilhado, ParticipanteRateio, RegraRateio
from classifier import ClassificadorFaturas
from integrador_planilha import IntegradorPlanilha
from pipeline_integracao import PipelineIntegracao, STATUS_ERRO_LEITURA, STATUS_REVISAO, STATUS_SUCESSO
from robo_vigia import MonitorPasta


def _config(tmp_path):
    valores = {
        'entrada_path': tmp_path / 'entrada', 'processamento_path': tmp_path / 'processando',
        'processadas_path': tmp_path / 'processadas', 'compartilhadas_path': tmp_path / 'compartilhadas',
        'exclusivas_path': tmp_path / 'exclusivas', 'revisao_path': tmp_path / 'revisao',
        'duplicadas_path': tmp_path / 'duplicadas', 'erro_path': tmp_path / 'erro',
        'backup_path': tmp_path / 'backups', 'planilha_oficial_path': tmp_path / 'copia.xlsx', 'ambiente': 'desenvolvimento',
    }
    return {chave: str(valor) for chave, valor in valores.items()}


def _planilha(config):
    workbook = Workbook()
    planilha = workbook.active
    planilha.append(['Serviço', 'Fornecedor', 'Complemento', 'Seção', 'Competência', 'SRA', 'DRF'])
    planilha.append(['Água e esgoto', 'CESAN', 'EDIFÍCIO-SEDE', 'SERVIÇOS DESPESAS COMPARTILHADAS', '05/2026', None, None])
    workbook.save(config['planilha_oficial_path'])
    workbook.close()


def _extrator(_path):
    return {'status': 'ok', 'faturas': [{'fornecedor': 'CESAN', 'servico': 'Água e esgoto', 'complemento': 'EDIFÍCIO-SEDE', 'secao': 'SERVIÇOS DESPESAS COMPARTILHADAS', 'competencia': '05/2026', 'valor': 100.0, 'descricao': 'edifício sede'}]}


def _regra(_dados):
    return RegraRateio('CESAN edifício-sede', (ParticipanteRateio('SRA', percentual=60), ParticipanteRateio('DRF', percentual=40)), origem='amostra autorizada')


def test_fluxo_completo_processa_uma_fatura_e_produz_relatorio(tmp_path):
    config = _config(tmp_path)
    for chave in ('entrada_path', 'processamento_path', 'processadas_path', 'revisao_path', 'erro_path', 'backup_path'):
        Path(config[chave]).mkdir(parents=True, exist_ok=True)
    _planilha(config)
    arquivo = Path(config['entrada_path']) / 'amostra_autorizada.pdf'
    arquivo.write_bytes(b'amostra autorizada')
    pipeline = PipelineIntegracao(config, extrator=_extrator, regra_compartilhada=_regra, classificador=ClassificadorFaturas(), calculador_compartilhado=CalculadorRateioCompartilhado(), integrador=IntegradorPlanilha(config), monitor=MonitorPasta(config, sleep=lambda _: None), enriquecer=lambda dados: dados)

    resultado = pipeline.processar_arquivo(arquivo)

    assert resultado['status'] == STATUS_SUCESSO
    assert [etapa['etapa'] for etapa in resultado['etapas']] == ['monitoramento', 'extracao', 'dados_para_classificacao', 'classificacao', 'calculo', 'validacao', 'planilha_auditoria', 'pasta_final']
    assert Path(resultado['arquivo_final']).parent == Path(config['processadas_path'])
    workbook = load_workbook(config['planilha_oficial_path'], data_only=False)
    assert workbook.active['F2'].value == 60
    assert workbook.active['G2'].value == 40


def test_pendencia_interrompe_planilha_e_nao_declara_sucesso_geral(tmp_path):
    config = _config(tmp_path)
    for chave in ('entrada_path', 'processamento_path', 'processadas_path', 'revisao_path', 'erro_path', 'backup_path'):
        Path(config[chave]).mkdir(parents=True, exist_ok=True)
    _planilha(config)
    arquivo = Path(config['entrada_path']) / 'amostra_revisao.pdf'
    arquivo.write_bytes(b'amostra')
    pipeline = PipelineIntegracao(config, extrator=lambda _path: {'status': 'ok', 'faturas': [{'fornecedor': 'CESAN', 'servico': 'Água e esgoto', 'valor': 100.0}]}, monitor=MonitorPasta(config, sleep=lambda _: None), integrador=IntegradorPlanilha(config))

    resultado = pipeline.executar_pasta()

    assert resultado['status_geral'] == 'PENDENTE'
    assert resultado['faturas'][0]['status'] == STATUS_REVISAO
    assert load_workbook(config['planilha_oficial_path']).active['F2'].value is None


def test_erro_de_leitura_vai_para_erro_e_nao_grava(tmp_path):
    config = _config(tmp_path)
    for chave in ('entrada_path', 'processamento_path', 'processadas_path', 'revisao_path', 'erro_path', 'backup_path'):
        Path(config[chave]).mkdir(parents=True, exist_ok=True)
    _planilha(config)
    arquivo = Path(config['entrada_path']) / 'ilegivel.pdf'
    arquivo.write_bytes(b'amostra')
    pipeline = PipelineIntegracao(config, extrator=lambda _path: {'status': 'ilegivel', 'faturas': [], 'erro': 'PDF ilegível'}, monitor=MonitorPasta(config, sleep=lambda _: None), integrador=IntegradorPlanilha(config))

    resultado = pipeline.processar_arquivo(arquivo)

    assert resultado['status'] == STATUS_ERRO_LEITURA
    assert Path(resultado['arquivo_final']).parent == Path(config['erro_path'])


def test_duplicidade_interrompe_antes_da_extracao(tmp_path):
    config = _config(tmp_path)
    for chave in ('entrada_path', 'processamento_path', 'processadas_path', 'revisao_path', 'erro_path', 'backup_path'):
        Path(config[chave]).mkdir(parents=True, exist_ok=True)
    _planilha(config)
    arquivo = Path(config['entrada_path']) / 'duplicada.pdf'
    arquivo.write_bytes(b'duplicada')
    banco = tmp_path / 'monitor.db'
    conn = sqlite3.connect(banco)
    conn.execute('CREATE TABLE faturas (hash_arquivo TEXT UNIQUE NOT NULL)')
    conn.execute('CREATE TABLE logs_processamento (fatura_id INTEGER, status TEXT, mensagem TEXT, hash_arquivo TEXT)')
    conn.commit()
    monitor = MonitorPasta(config, connection_factory=lambda: sqlite3.connect(banco), sleep=lambda _: None)
    conn.execute('INSERT INTO faturas (hash_arquivo) VALUES (?)', (monitor.calcular_hash(arquivo),))
    conn.commit()
    conn.close()
    chamadas = []
    pipeline = PipelineIntegracao(config, extrator=lambda path: chamadas.append(path), monitor=monitor, integrador=IntegradorPlanilha(config))

    resultado = pipeline.processar_arquivo(arquivo)

    assert resultado['status'] == 'DUPLICADA_OU_JÁ_PROCESSADA'
    assert chamadas == []
    assert Path(resultado['arquivo_final']).parent == Path(config['duplicadas_path'])


def test_erro_de_gravacao_nao_vai_para_processadas(tmp_path):
    config = _config(tmp_path)
    for chave in ('entrada_path', 'processamento_path', 'processadas_path', 'revisao_path', 'erro_path', 'backup_path'):
        Path(config[chave]).mkdir(parents=True, exist_ok=True)
    _planilha(config)
    arquivo = Path(config['entrada_path']) / 'erro_gravacao.pdf'
    arquivo.write_bytes(b'amostra')

    class IntegradorComErro:
        def gravar(self, fatura, calculo):
            return {'status': 'ERRO_DE_GRAVAÇÃO', 'motivos_revisao': ['Planilha bloqueada.']}

    pipeline = PipelineIntegracao(config, extrator=_extrator, regra_compartilhada=_regra, monitor=MonitorPasta(config, sleep=lambda _: None), integrador=IntegradorComErro())
    resultado = pipeline.processar_arquivo(arquivo)

    assert resultado['status'] == 'ERRO_DE_GRAVAÇÃO'
    assert Path(resultado['arquivo_final']).parent == Path(config['erro_path'])
    assert not list(Path(config['processadas_path']).iterdir())