from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from integrador_planilha import IntegradorPlanilha


def _planilha(tmp_path):
    caminho = tmp_path / 'copia_teste.xlsx'
    workbook = Workbook()
    planilha = workbook.active
    planilha.append(['Serviço', 'Fornecedor', 'Complemento', 'Seção', 'Competência', 'SRA', 'DRF', 'CGU', 'Total'])
    planilha.append(['Água e esgoto', 'CESAN', 'ESTACIONAMENTO', 'SERVIÇOS EXCLUSIVOS', '05/2026', None, 0, 0, '=SUM(F2:H2)'])
    planilha['I2'].font = Font(bold=True)
    workbook.save(caminho)
    workbook.close()
    return caminho


def _config(caminho, tmp_path):
    return {
        'planilha_oficial_path': str(caminho),
        'backup_path': str(tmp_path / 'backups'),
        'ambiente': 'desenvolvimento',
    }


def _calculo(parcelas):
    return {'status': 'CALCULADO_VALIDADO', 'regra': {'nome': 'regra teste'}, 'parcelas': parcelas}


def _fatura():
    return {
        'servico': 'Água e esgoto',
        'fornecedor': 'CESAN',
        'complemento': 'ESTACIONAMENTO',
        'secao': 'SERVIÇOS EXCLUSIVOS',
        'competencia': '05/2026',
    }


def test_grava_na_copia_preserva_formula_estilo_backup_e_auditoria(tmp_path):
    caminho = _planilha(tmp_path)
    auditoria = []
    resultado = IntegradorPlanilha(_config(caminho, tmp_path), usuario='teste', audit_sink=auditoria.append).gravar(
        _fatura(), _calculo([
            {'orgao': 'SRA', 'valor': '100.00'},
            {'orgao': 'DRF', 'valor': '0.00'},
            {'orgao': 'CGU', 'valor': '0.00'},
        ])
    )

    assert resultado['status'] == 'GRAVADO_VALIDADO'
    assert len(list((tmp_path / 'backups').glob('*.xlsx'))) == 1
    workbook = load_workbook(caminho, data_only=False)
    planilha = workbook.active
    assert planilha['F2'].value == 100
    assert planilha['G2'].value == 0
    assert planilha['I2'].value == '=SUM(F2:H2)'
    assert planilha['I2'].font.bold is True
    assert auditoria[0]['celula'] == 'F2'
    assert auditoria[0]['valor_anterior'] is None
    assert auditoria[0]['valor_novo'] == 100.0
    assert auditoria[0]['usuario'] == 'teste'
    assert auditoria[0]['regra']['nome'] == 'regra teste'


def test_linha_inexistente_vai_para_revisao_sem_criar_linha(tmp_path):
    caminho = _planilha(tmp_path)
    fatura = _fatura()
    fatura['competencia'] = '06/2026'
    resultado = IntegradorPlanilha(_config(caminho, tmp_path)).gravar(fatura, _calculo([{'orgao': 'SRA', 'valor': '100.00'}]))

    assert resultado['status'] == 'AGUARDANDO_REVISÃO'
    assert 'não encontrada' in resultado['motivos_revisao'][0]
    assert load_workbook(caminho).active.max_row == 2


def test_configuracao_ausente_bloqueia_gravacao(tmp_path):
    caminho = _planilha(tmp_path)
    config = _config(caminho, tmp_path)
    config['planilha_oficial_path'] = ''
    resultado = IntegradorPlanilha(config).gravar(_fatura(), _calculo([{'orgao': 'SRA', 'valor': '100.00'}]))

    assert resultado['status'] == 'AGUARDANDO_REVISÃO'
    assert 'não configurado' in resultado['motivos_revisao'][0]


def test_ambiente_producao_bloqueia_teste(tmp_path):
    caminho = _planilha(tmp_path)
    config = _config(caminho, tmp_path)
    config['ambiente'] = 'producao'
    resultado = IntegradorPlanilha(config).gravar(_fatura(), _calculo([{'orgao': 'SRA', 'valor': '100.00'}]))

    assert resultado['status'] == 'AGUARDANDO_REVISÃO'
    assert 'produção' in resultado['motivos_revisao'][0]